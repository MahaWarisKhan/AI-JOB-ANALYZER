import re
from jobs import recommend_jobs
from resume_parser import parse_resume_details
from flask import Flask, render_template, request
from pypdf import PdfReader
from docx import Document

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

SKILLS = {
    "Python": ["python"],
    "Java": ["java"],
    "JavaScript": ["javascript"],
    "HTML": ["html", "html5"],
    "CSS": ["css", "css3"],
    "SQL": ["sql"],
    "MySQL": ["mysql"],
    "SQL Server": ["sql server", "mssql"],
    "Flask": ["flask"],
    "Django": ["django"],
    "React": ["react", "reactjs", "react.js"],
    "Git": ["git"],
    "Machine Learning": ["machine learning"],
    "NLP": ["nlp", "natural language processing"],
    "Excel": ["excel", "microsoft excel"],
}


def extract_text(file):
    if file.filename.lower().endswith(".pdf"):
        reader = PdfReader(file.stream)
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        )

    document = Document(file.stream)
    parts = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)


def find_skills(text):
    found = []

    for skill, aliases in SKILLS.items():
        for alias in aliases:
            pattern = r"(?<!\w)" + re.escape(alias) + r"(?!\w)"
            if re.search(pattern, text, flags=re.IGNORECASE):
                found.append(skill)
                break

    return sorted(found)


@app.route("/", methods=["GET", "POST"])
def dashboard():
    text = ""
    skills = []
    filename = ""
    error = ""

    if request.method == "POST":
        file = request.files.get("resume")

        if not file or not file.filename:
            error = "Please select your resume."

        elif not file.filename.lower().endswith((".pdf", ".docx")):
            error = "Please upload a PDF or DOCX file."

        else:
            try:
                text = extract_text(file).strip()

                if not text:
                    error = (
                        "No readable text found. Please upload a "
                        "text-based resume. Scanned PDFs need OCR."
                    )
                else:
                    filename = file.filename
                    skills = find_skills(text)

            except Exception:
                text = ""
                error = (
                    "We could not read this file. Please upload "
                    "a valid, unprotected PDF or DOCX."
                )

    recommendations = recommend_jobs(skills) if text else []
    details = parse_resume_details(text)

    return render_template(
        "dashboard.html",
        text=text,
        skills=skills,
        filename=filename,
        error=error,
        recommendations=recommendations,
        education=details["education"],
        experience=details["experience"],
    )


@app.errorhandler(413)
def too_large(error):
    return render_template(
        "dashboard.html",
        text="",
        skills=[],
        filename="",
        error="Your file exceeds the 5 MB upload limit.",
    ), 413


if __name__ == "__main__":
    app.run()